from flask import Flask, request, jsonify
from flask_sqlalchemy import SQLAlchemy
from flask_cors import CORS
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import cm
from flask import send_file
from datetime import date
from sqlalchemy import func
from openpyxl import Workbook

import io
import os
from dotenv import load_dotenv

load_dotenv()

app = Flask(__name__)

frontend_url = os.getenv(
    "FRONTEND_URL",
    "http://localhost:5173"
)

CORS(
    app,
    resources={
        r"/*": {
            "origins": [
                "http://localhost:5173",
                frontend_url
            ]
        }
    }
)

database_url = os.getenv("DATABASE_URL")

if database_url and database_url.startswith("postgresql://"):
    database_url = database_url.replace(
        "postgresql://",
        "postgresql+psycopg2://",
        1
    )

app.config["SQLALCHEMY_DATABASE_URI"] = database_url

app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False


db = SQLAlchemy(app)

class Stock(db.Model):
    __tablename__ = "stocks"

    id = db.Column(db.Integer, primary_key=True)
    product_name = db.Column(db.String(100), nullable=False)
    quantity = db.Column(db.Integer, nullable=False)
    cost_price = db.Column(db.Float, nullable=False)
    selling_price = db.Column(db.Float, nullable=False)
    created_at = db.Column(db.DateTime, server_default=db.func.now())

class Sale(db.Model):
    __tablename__ = "sales"

    id = db.Column(db.Integer, primary_key=True)
    stock_id = db.Column(db.Integer, db.ForeignKey("stocks.id"), nullable=False)
    quantity_sold = db.Column(db.Integer, nullable=False)
    selling_price = db.Column(db.Float, nullable=False)
    sold_at = db.Column(db.DateTime, server_default=db.func.now())

    stock = db.relationship("Stock")
    
def stock_to_dict(stock):
    profit_per_unit = stock.selling_price - stock.cost_price
    total_profit = profit_per_unit * stock.quantity

    return {
        "id": stock.id,
        "product_name": stock.product_name,
        "quantity": stock.quantity,
        "cost_price": stock.cost_price,
        "selling_price": stock.selling_price,
        "profit_per_unit": profit_per_unit,
        "total_profit": total_profit,
        "created_at": stock.created_at
    }

@app.route("/")
def home():
    return jsonify({
        "app": "StockThozhan",
        "status": "running"
    })
    
@app.route("/add-stock", methods=["POST"])
def add_stock():
    data = request.json

    if not data.get("product_name", "").strip():
        return jsonify({"error": "Product name is required"}), 400

    if data["quantity"] <= 0:
        return jsonify({"error": "Quantity must be greater than 0"}), 400

    if data["cost_price"] <= 0 or data["selling_price"] <= 0:
        return jsonify({"error": "Prices must be greater than 0"}), 400

    new_stock = Stock(
        product_name=data["product_name"].strip(),
        quantity=data["quantity"],
        cost_price=data["cost_price"],
        selling_price=data["selling_price"]
    )

    db.session.add(new_stock)
    db.session.commit()

    return jsonify({"message": "Stock added successfully"}), 201

@app.route("/stocks", methods=["GET"])
def get_stocks():
    stocks = Stock.query.order_by(Stock.id.asc()).all()
    result = [stock_to_dict(stock) for stock in stocks]
    return jsonify(result), 200

@app.route("/delete-stock/<int:stock_id>", methods=["DELETE"])
def delete_stock(stock_id):
    stock = db.session.get(Stock, stock_id)

    if not stock:
        return jsonify({"error": "Stock not found"}), 404

    sale_exists = db.session.execute(
        db.text("SELECT 1 FROM sales WHERE stock_id = :id LIMIT 1"),
        {"id": stock_id}
    ).first()

    if sale_exists:
        return jsonify({
            "error": "Cannot delete product. Sales already recorded for this item."
        }), 400

    db.session.delete(stock)
    db.session.commit()

    return jsonify({"message": "Stock deleted successfully"}), 200


@app.route("/update-stock/<int:stock_id>", methods=["PUT"])
def update_stock(stock_id):
    stock = db.session.get(Stock, stock_id)

    if not stock:
        return jsonify({"error": "Stock not found"}), 404

    data = request.json

    if not data.get("product_name", "").strip():
        return jsonify({"error": "Product name is required"}), 400

    if data["quantity"] <= 0:
        return jsonify({"error": "Quantity must be greater than 0"}), 400

    if data["cost_price"] <= 0 or data["selling_price"] <= 0:
        return jsonify({"error": "Prices must be greater than 0"}), 400

    stock.product_name = data["product_name"].strip()
    stock.quantity = data["quantity"]
    stock.cost_price = data["cost_price"]
    stock.selling_price = data["selling_price"]

    db.session.commit()

    return jsonify({"message": "Stock updated successfully"}), 200

@app.route("/add-sale", methods=["POST"])
def add_sale():
    data = request.json

    stock = db.session.get(Stock, data.get("stock_id"))
    if not stock:
        return jsonify({"error": "Stock not found"}), 404

    qty = int(data.get("quantity_sold", 0))
    if qty <= 0:
        return jsonify({"error": "Invalid quantity"}), 400

    if qty > stock.quantity:
        return jsonify({"error": "Not enough stock"}), 400

    sale = Sale(
        stock_id=stock.id,
        quantity_sold=qty,
        selling_price=stock.selling_price
    )

    stock.quantity -= qty   

    db.session.add(sale)
    db.session.commit()

    return jsonify({"message": "Sale recorded"}), 201
    

@app.route("/daily-sales", methods=["GET"])
def daily_sales():
    today = date.today()

    rows = (
        db.session.query(
            Stock.product_name,
            func.sum(Sale.quantity_sold).label("total_sold"),
            func.sum(Sale.quantity_sold * Sale.selling_price).label("total_amount")
        )
        .join(Sale, Sale.stock_id == Stock.id)
        .filter(func.date(Sale.sold_at) == today)
        .group_by(Stock.product_name)
        .all()
    )

    return jsonify([
        {
            "product": r.product_name,
            "quantity": r.total_sold,
            "amount": r.total_amount
        }
        for r in rows
    ])

@app.route("/sales-history", methods=["GET"])
def sales_history():
    start = request.args.get("start")
    end = request.args.get("end")

    query = (
        db.session.query(
            Sale.sold_at,
            Stock.product_name,
            Sale.quantity_sold,
            (Sale.quantity_sold * Sale.selling_price).label("total_amount")
        )
        .join(Stock)
    )

    if start and end:
        query = query.filter(
            db.func.date(Sale.sold_at).between(start, end)
        )

    results = query.order_by(Sale.sold_at.desc()).all()

    return jsonify([
        {
            "date": r.sold_at.strftime("%Y-%m-%d"),
            "product": r.product_name,
            "quantity": r.quantity_sold,
            "amount": r.total_amount
        }
        for r in results
    ])

@app.route("/sales-history/excel", methods=["GET"])
def sales_history_excel():
    start = request.args.get("start")
    end = request.args.get("end")

    query = (
        db.session.query(
            Sale.sold_at,
            Stock.product_name,
            Sale.quantity_sold,
            (Sale.quantity_sold * Sale.selling_price).label("amount")
        )
        .join(Stock)
    )

    if start and end:
        query = query.filter(
            func.date(Sale.sold_at).between(start, end)
        )

    results = query.order_by(Sale.sold_at.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales History"

    ws.append(["Date", "Product", "Quantity Sold", "Total Amount"])

    for r in results:
        ws.append([
            r.sold_at.strftime("%Y-%m-%d"),
            r.product_name,
            r.quantity_sold,
            r.amount
        ])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    return send_file(
        stream,
        as_attachment=True,
        download_name="sales_history.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )



@app.route("/download-excel", methods=["GET"])
def download_excel():
    stocks = Stock.query.all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Report"
    
    ws.append([
        "Product Name",
        "Quantity",
        "Cost Price",
        "Selling Price",
        "Profit / Unit",
        "Total Profit"
    ])

    for stock in stocks:
        profit_per_unit = stock.selling_price - stock.cost_price
        total_profit = profit_per_unit * stock.quantity

        ws.append([
            stock.product_name,
            stock.quantity,
            stock.cost_price,
            stock.selling_price,
            profit_per_unit,
            total_profit
        ])

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    return send_file(
        file_stream,
        as_attachment=True,
        download_name="stockthozhan_report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route("/daily-sales-excel", methods=["GET"])
def daily_sales_excel():
    today = date.today()

    sales = (
        db.session.query(
            Stock.product_name,
            db.func.sum(Sale.quantity_sold).label("quantity"),
            db.func.sum(Sale.quantity_sold * Sale.selling_price).label("amount")
        )
        .join(Stock, Stock.id == Sale.stock_id)
        .filter(db.func.date(Sale.sold_at) == today)
        .group_by(Stock.product_name)
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Sales"

    ws.append(["Product", "Quantity Sold", "Total Amount"])

    for s in sales:
        ws.append([s.product_name, s.quantity, s.amount])

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    return send_file(
        stream,
        as_attachment=True,
        download_name=f"daily_sales_{today}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route("/download-daily-sales-pdf", methods=["GET"])
def download_daily_sales_pdf():
    today = date.today()

    results = (
        db.session.query(
            Stock.product_name,
            func.sum(Sale.quantity_sold).label("quantity"),
            func.sum(Sale.quantity_sold * Sale.selling_price).label("amount")
        )
        .join(Stock, Stock.id == Sale.stock_id)
        .filter(func.date(Sale.sold_at) == today)
        .group_by(Stock.product_name)
        .all()
    )

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    _, height = A4

    y = height - 2 * cm

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(2 * cm, y, "StockThozhan - Daily Sales Report")
    y -= 1 * cm

    pdf.setFont("Helvetica", 11)
    pdf.drawString(2 * cm, y, f"Date: {today}")
    y -= 1 * cm

    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(2 * cm, y, "Product")
    pdf.drawString(8 * cm, y, "Quantity")
    pdf.drawString(12 * cm, y, "Amount")
    y -= 0.7 * cm

    pdf.setFont("Helvetica", 10)

    total_amount = 0

    for r in results:
        pdf.drawString(2 * cm, y, r.product_name)
        pdf.drawString(8 * cm, y, str(int(r.quantity)))
        pdf.drawString(12 * cm, y, f"₹ {float(r.amount)}")

        total_amount += float(r.amount)
        y -= 0.6 * cm

        if y < 2 * cm:
            pdf.showPage()
            y = height - 2 * cm
            pdf.setFont("Helvetica", 10)

    y -= 1 * cm
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(2 * cm, y, f"Total Sales Amount: ₹ {total_amount}")

    pdf.save()
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"daily_sales_{today}.pdf",
        mimetype="application/pdf"
    )


@app.route("/download-pdf", methods=["GET"])
def download_pdf():
    stocks = Stock.query.all()

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    _, height = A4

    y = height - 2 * cm

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(2 * cm, y, "StockThozhan - Stock Report")
    y -= 1 * cm

    pdf.setFont("Helvetica", 11)
    pdf.drawString(2 * cm, y, "-------------------------------------------")
    y -= 1 * cm

    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(2 * cm, y, "Product")
    pdf.drawString(6 * cm, y, "Qty")
    pdf.drawString(8 * cm, y, "Cost")
    pdf.drawString(10 * cm, y, "Selling")
    pdf.drawString(13 * cm, y, "Profit")
    y -= 0.7 * cm

    pdf.setFont("Helvetica", 10)

    total_profit = 0

    for stock in stocks:
        profit_per_unit = stock.selling_price - stock.cost_price
        row_profit = profit_per_unit * stock.quantity
        total_profit += row_profit

        pdf.drawString(2 * cm, y, stock.product_name)
        pdf.drawString(6 * cm, y, str(stock.quantity))
        pdf.drawString(8 * cm, y, str(stock.cost_price))
        pdf.drawString(10 * cm, y, str(stock.selling_price))
        pdf.drawString(13 * cm, y, str(row_profit))

        y -= 0.6 * cm

        if y < 2 * cm:
            pdf.showPage()
            y = height - 2 * cm
            pdf.setFont("Helvetica", 10)

    y -= 1 * cm
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(2 * cm, y, f"Overall Profit: ₹ {total_profit}")

    pdf.save()
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="stockthozhan_report.pdf",
        mimetype="application/pdf"
    )

@app.route("/sales-history/pdf", methods=["GET"])
def sales_history_pdf():
    start = request.args.get("start")
    end = request.args.get("end")

    query = (
        db.session.query(
            Sale.sold_at,
            Stock.product_name,
            Sale.quantity_sold,
            (Sale.quantity_sold * Sale.selling_price).label("amount")
        )
        .join(Stock)
    )

    if start and end:
        query = query.filter(
            func.date(Sale.sold_at).between(start, end)
        )

    results = query.order_by(Sale.sold_at.asc()).all()

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    _, height = A4

    y = height - 2 * cm

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(2 * cm, y, "Sales History Report")
    y -= 1 * cm

    if start and end:
        pdf.setFont("Helvetica", 11)
        pdf.drawString(2 * cm, y, f"From {start} to {end}")
        y -= 1 * cm

    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(2 * cm, y, "Date")
    pdf.drawString(5 * cm, y, "Product")
    pdf.drawString(10 * cm, y, "Qty")
    pdf.drawString(13 * cm, y, "Amount")
    y -= 0.7 * cm

    pdf.setFont("Helvetica", 10)
    total_amount = 0

    for r in results:
        pdf.drawString(2 * cm, y, r.sold_at.strftime("%Y-%m-%d"))
        pdf.drawString(5 * cm, y, r.product_name)
        pdf.drawString(10 * cm, y, str(r.quantity_sold))
        pdf.drawString(13 * cm, y, f"₹ {r.amount}")

        total_amount += r.amount
        y -= 0.6 * cm

        if y < 2 * cm:
            pdf.showPage()
            y = height - 2 * cm
            pdf.setFont("Helvetica", 10)

    y -= 1 * cm
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(2 * cm, y, f"Total Amount: ₹ {total_amount}")

    pdf.save()
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="sales_history.pdf",
        mimetype="application/pdf"
    )

@app.route("/reset-all", methods=["DELETE"])
def reset_all():
    try:
        db.session.execute(db.text("DELETE FROM sales"))
        db.session.execute(db.text("DELETE FROM stocks"))
        db.session.commit()

        return jsonify({"message": "All data reset successfully"}), 200

    except Exception as e:
        db.session.rollback()
        return jsonify({"error": str(e)}), 500

with app.app_context():
    db.create_all()

if __name__ == "__main__":
    app.run(debug=False)
